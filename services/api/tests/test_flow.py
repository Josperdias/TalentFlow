"""End-to-end API integration tests using isolated per-suite DB."""
import os,tempfile
from pathlib import Path
from fastapi.testclient import TestClient

DB=Path(tempfile.gettempdir())/'talentflow_test_suite_20261011.db'
try:DB.unlink()
except FileNotFoundError:pass
os.environ['DATABASE_URL']=f'sqlite:///{DB}'
from app.main import app


def auth(client,email='admin@nexa.demo',password='Demo@2026!'):
    r=client.post('/api/auth/login',json={'email':email,'password':password})
    assert r.status_code==200,r.text
    csrf=client.get('/api/auth/csrf')
    assert csrf.status_code==200,csrf.text
    client.headers.update({'X-TF-CSRF':csrf.json()['token']})
    return r.json()['user']

def test_health_login_and_seed():
    with TestClient(app) as c:
        assert c.get('/api/health').json()['status']=='ok'
        assert c.get('/api/me').status_code==401
        assert c.post('/api/auth/login',json={'email':'admin@nexa.demo','password':'wrong'}).status_code==401
        auth(c)
        assert len(c.get('/api/jobs').json())>=10
        assert len(c.get('/api/candidates').json())>=100
        m=c.get('/api/dashboard').json()['metrics']
        assert m['candidates_active']>0
        assert m['jobs_open']>0
        assert m['interviews']>0
        assert m['hired']>0
        assert c.get('/').status_code==200

def test_full_recruiting_lifecycle():
    with TestClient(app) as c:
        auth(c)
        job=c.post('/api/jobs',json={'title':'Analista de Gente e Gestão','department':'Pessoas','status':'open','requirements':'Excel recrutamento comunicação','description':'Atuar na área de RH'});assert job.status_code==200,job.text
        jobid=job.json()['id']
        candidate=c.post('/api/candidates',json={'name':'Teste Integral','email':'teste.integral@example.com','skills':'Excel, Comunicação, Recrutamento','experience_years':4});assert candidate.status_code==200,candidate.text
        cid=candidate.json()['id']
        a=c.post('/api/applications',json={'candidate_id':cid,'job_id':jobid});assert a.status_code==200,a.text
        aid=a.json()['id']
        assert c.post('/api/applications',json={'candidate_id':cid,'job_id':jobid}).status_code==409
        stages=c.get(f'/api/jobs/{jobid}').json()['stages'];assert len(stages)==9
        s=stages[1]['id'];m=c.post(f'/api/applications/{aid}/move',json={'stage_id':s});assert m.status_code==200,m.text
        appdata=c.get(f'/api/applications/{aid}').json();assert appdata['application']['stage']=='Triagem'
        assert any(e['kind']=='stage_changed' for e in appdata['activities'])
        score=c.post(f'/api/applications/{aid}/fit',json={});assert score.status_code==200,score.text
        manual=c.post(f'/api/applications/{aid}/fit',json={'score':88,'reason':'Revisado manualmente com evidências de experiência em Excel'});assert manual.status_code==200,manual.text
        it=c.post(f'/api/applications/{aid}/interviews',json={'starts_at':'2026-12-01T14:30','duration_minutes':45});assert it.status_code==200,it.text
        assert c.patch('/api/interviews/'+str(it.json()['id']),json={'status':'confirmed'}).status_code==200
        sc=c.post(f'/api/applications/{aid}/scorecards',json={'criteria':[{'name':'Excel','score':5,'weight':2},{'name':'Comunicação','score':3,'weight':1}],'notes':'Aprovado tecnicamente'});assert sc.status_code==200,sc.text
        assert sc.json()['result']==4.33
        comm=c.post(f'/api/applications/{aid}/communications',json={'subject':'Olá {{nome}}','body':'Candidatura {{vaga}}','channel':'email'});assert comm.status_code==200,comm.text
        assert c.get(f'/api/applications/{aid}').json()['communications'][0]['status']=='queued_mock'
        offer=c.post(f'/api/applications/{aid}/offers',json={'salary':6750,'start_date':'2026-12-10'});assert offer.status_code==200,offer.text
        oid=offer.json()['id'];sent=c.post(f'/api/offers/{oid}/send');assert sent.status_code==200,sent.text
        token=sent.json()['accept_path'].split('/')[-1]
        assert c.get(f'/api/public/offer/{token}').json()['salary']==6750
        assert c.post(f'/api/public/offer/{token}/respond',json={'response':'accepted'}).json()['status']=='accepted'
        assert c.post(f'/api/public/offer/{token}/respond',json={'response':'accepted'}).status_code==409
        hired=next(s for s in stages if s['kind']=='hired')
        assert c.post(f'/api/applications/{aid}/move',json={'stage_id':hired['id']}).status_code==422
        hired_result=c.post(f'/api/applications/{aid}/move',json={'stage_id':hired['id'],'reason':'Proposta aceita e contratação formalizada'});assert hired_result.status_code==200,hired_result.text
        assert hired_result.json()['status']=='hired'
        assert c.get('/api/audit').status_code==200
        assert any(e['action']=='stage_changed' for e in c.get('/api/audit').json())

def test_careers_public_apply_and_lgpd():
    with TestClient(app) as c:
        jobs=c.get('/api/public/jobs').json();assert len(jobs)>0
        public_job=jobs[0]['id']
        invalid=c.post('/api/public/apply',json={'job_id':public_job,'name':'Maria de Teste','email':'maria@example.com','consent':False});assert invalid.status_code==422
        valid=c.post('/api/public/apply',json={'job_id':public_job,'name':'Maria de Teste','email':'maria@example.com','consent':True,'resume_text':'Testes e processos digitais'});assert valid.status_code==200,valid.text
        auth(c)
        assert any(x['email']=='maria@example.com' for x in c.get('/api/candidates').json())

def test_permissions_and_tenant_isolation():
    with TestClient(app) as c:
        auth(c,'admin@orbe.demo')
        assert c.get('/api/candidates').json()==[]
        assert c.get('/api/jobs').json()==[]
        assert c.get('/api/jobs/1').status_code==404
        assert c.get('/api/applications/1').status_code==404
        assert c.post('/api/applications/1/move',json={'stage_id':1}).status_code==404
        assert c.get('/api/dashboard').json()['metrics']['hired']==0
    with TestClient(app) as c:
        auth(c,'entrevistador@nexa.demo')
        assert c.post('/api/jobs',json={'title':'Não permitido'}).status_code==403
        assert c.post('/api/candidates',json={'name':'Não permitido','email':'x@example.com'}).status_code==403
        assert c.get('/api/audit').status_code==403
        assert c.get('/api/jobs').status_code==200
    with TestClient(app) as c:
        auth(c,'candidato@nexa.demo')
        assert c.get('/api/jobs').status_code==403
        assert c.get('/api/candidates').status_code==403


def test_automation_and_tasks():
    with TestClient(app) as c:
        auth(c)
        rules=c.get('/api/automations').json();assert len(rules)>=2
        r=rules[0];toggle=c.post(f'/api/automations/{r["id"]}/toggle');assert toggle.status_code==200
        assert toggle.json()['enabled']==False
        assert c.post(f'/api/automations/{r["id"]}/toggle').json()['enabled']==True
        tasks=c.get('/api/tasks').json();assert len(tasks)>=5
        assert c.post(f'/api/tasks/{tasks[0]["id"]}/toggle').status_code==200

def test_resume_parsing_and_secure_download():
    from io import BytesIO
    from docx import Document
    with TestClient(app) as c:
        auth(c)
        candidate=c.post('/api/candidates',json={'name':'Curriculo Teste','email':'curriculo@example.com'}).json()
        did=candidate['id']
        doc=Document();doc.add_paragraph('Curriculo Teste');doc.add_paragraph('Excel, Recrutamento e Python');doc.add_paragraph('curriculo@example.com')
        file=BytesIO();doc.save(file)
        r=c.post(f'/api/candidates/{did}/resume',files={'file':('cv.docx',file.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        assert r.status_code==200,r.text
        assert r.json()['extracted_chars']>20
        resume_id=r.json()['id']
        assert c.get(f'/api/resumes/{resume_id}/download').status_code==200
        assert 'Python' in c.get(f'/api/candidates/{did}').json()['candidate']['resume_text']
        assert c.post(f'/api/candidates/{did}/resume',files={'file':('cv.exe',b'malicious','application/octet-stream')}).status_code==422
    with TestClient(app) as second:
        auth(second,'admin@orbe.demo')
        assert second.get(f'/api/resumes/{resume_id}/download').status_code==404


def test_job_approvals_and_anonymization():
    with TestClient(app) as c:
        auth(c)
        j=c.post('/api/jobs',json={'title':'Analista Legal','status':'approval'}).json()
        assert c.post(f'/api/jobs/{j["id"]}/approve',json={'decision':'approved'}).json()['status']=='open'
        assert c.post(f'/api/jobs/{j["id"]}/approve',json={'decision':'approved'}).status_code==409
        cand=c.post('/api/candidates',json={'name':'Perfil para Exclusão','email':'exclusao@example.com'}).json()
        assert c.post(f'/api/candidates/{cand["id"]}/anonymize').status_code==200
        assert all(ca['id']!=cand['id'] for ca in c.get('/api/candidates').json())


def test_candidate_private_portal_and_assignments():
    with TestClient(app) as c:
        auth(c,'candidato@nexa.demo')
        items=c.get('/api/portal');assert items.status_code==200,items.text
        own=items.json()['applications'];assert len(own)>=1
        assert all('job' in a and 'stage' in a for a in own)
        assert c.get('/api/candidates').status_code==403
        assert c.get('/api/jobs').status_code==403
    with TestClient(app) as other:
        auth(other,'admin@orbe.demo')
        assert other.get('/api/portal').status_code==403
        assert other.get('/api/interviews').json()==[]
        assert other.get('/api/offers').json()==[]
        assert other.get('/api/tasks').json()==[]

def test_lifecycle_core_and_evidence_metrics():
    with TestClient(app) as c:
        auth(c)
        people=c.get('/api/lifecycle/employees'); assert people.status_code==200
        assert len(people.json())==24
        analytics=c.get('/api/lifecycle/analytics');assert analytics.status_code==200,analytics.text
        metrics=analytics.json()['indicators']
        assert metrics['active_headcount']['value']==22
        assert 0<=metrics['learning_completion']['value']<=100
        assert metrics['retention_90d']['denominator']>0
        emp=people.json()[0];pid=emp['id']
        detail=c.get(f'/api/lifecycle/employees/{pid}').json();assert detail['employee']['name']==emp['name']
        t=c.post(f'/api/lifecycle/employees/{pid}/onboarding',json={'title':'Treino de segurança','phase':'Primeira semana','due_date':'2026-12-20'});assert t.status_code==200,t.text
        tid=t.json()['id'];assert c.post(f'/api/lifecycle/onboarding/{tid}/toggle').json()['done'] is True
        assert c.post(f'/api/lifecycle/onboarding/{tid}/toggle').json()['done'] is False
        course=c.post('/api/lifecycle/courses',json={'title':'Novo curso de gestão de projetos','skill':'Planejamento','duration_hours':4});assert course.status_code==200,course.text
        course_id=course.json()['id']
        assigned=c.post('/api/lifecycle/enrollments',json={'employee_id':pid,'course_id':course_id});assert assigned.status_code==200,assigned.text
        assert c.post('/api/lifecycle/enrollments',json={'employee_id':pid,'course_id':course_id}).status_code==409
        assert c.post(f'/api/lifecycle/enrollments/{assigned.json()["id"]}/complete').json()['status']=='completed'
        goal=c.post(f'/api/lifecycle/employees/{pid}/goals',json={'title':'Concluir 10 entregas','target_value':10,'metric_unit':'entregas','due_date':'2026-12-30'});assert goal.status_code==200,goal.text
        progress=c.post(f'/api/lifecycle/goals/{goal.json()["id"]}/progress',json={'actual_value':11});assert progress.json()['status']=='achieved'
        review=c.post(f'/api/lifecycle/employees/{pid}/reviews',json={'period':'2026-Q4','criteria':[{'name':'Resultados','score':5,'weight':2},{'name':'Comunicação','score':3,'weight':1}],'feedback':'Evidências de desempenho registradas.'});assert review.status_code==200,review.text
        assert review.json()['score']==4.33
        assert c.post(f'/api/lifecycle/employees/{pid}/reviews',json={'period':'2026-Q4','criteria':[{'score':8,'weight':1}]}).status_code==422
        assert c.post(f'/api/lifecycle/employees/{pid}/career',json={'kind':'transfer','description':'Movimentação para equipe produto','effective_on':'2026-11-01'}).status_code==200
        audits=c.get('/api/audit').json();assert any(x['action']=='learning_completed' for x in audits)

def test_candidate_to_employee_and_cross_company_rejections():
    with TestClient(app) as c:
        auth(c)
        job=c.post('/api/jobs',json={'title':'Analista de Talentos - RH','department':'Pessoas','status':'open'}).json()
        cand=c.post('/api/candidates',json={'name':'Funcionário Contratado','email':'contratado@example.com','skills':'RH, Python','experience_years':3}).json()
        a=c.post('/api/applications',json={'candidate_id':cand['id'],'job_id':job['id']}).json()
        aid=a['id']
        assert c.post(f'/api/lifecycle/from-application/{aid}').status_code==409
        offer=c.post(f'/api/applications/{aid}/offers',json={'salary':6700,'start_date':'2026-12-20'}).json()
        c.post(f'/api/offers/{offer["id"]}/send')
        c.post('/api/public/offer/'+offer['accept_path'].split('/')[-1]+'/respond',json={'response':'accepted'})
        stage=next(s for s in c.get('/api/jobs/'+str(job['id'])).json()['stages'] if s['kind']=='hired')
        c.post(f'/api/applications/{aid}/move',json={'stage_id':stage['id'],'reason':'Oferta aceita e admissão confirmada'})
        r=c.post(f'/api/lifecycle/from-application/{aid}');assert r.status_code==200,r.text
        person=r.json();assert person['application_id']==aid
        assert len(c.get('/api/lifecycle/employees/'+str(person['id'])).json()['onboarding'])==5
        assert c.post(f'/api/lifecycle/from-application/{aid}').status_code==409
    with TestClient(app) as orbe:
        auth(orbe,'admin@orbe.demo')
        assert orbe.get('/api/lifecycle/employees').json()==[]
        assert orbe.get('/api/lifecycle/analytics').json()['indicators']['active_headcount']['value']==0
        assert orbe.get('/api/lifecycle/employees/'+str(person['id'])).status_code==404
        assert orbe.post('/api/lifecycle/enrollments',json={'employee_id':person['id'],'course_id':1}).status_code==404
        assert orbe.post(f'/api/lifecycle/from-application/{aid}').status_code==404
    with TestClient(app) as interviewer:
        auth(interviewer,'entrevistador@nexa.demo')
        assert interviewer.get('/api/lifecycle/employees').status_code==403
        assert interviewer.post('/api/lifecycle/employees',json={'name':'Não permitido','work_email':'x@example.test','hire_date':'2026-10-11'}).status_code==403
